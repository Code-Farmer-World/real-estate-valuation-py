"""在產出的 xlsx 裡加一張「計算依據」工作表。

官方書表只有數字，沒有地方寫「這一格為什麼是這個值」。審查或訴願時那些依據是
真正被問的東西，所以另外開一張表把它攤出來，與書表放在同一個檔案，
不會分家。

這一層不 import kernel。依據由呼叫端算好後以參數傳入（duck typing，
只依賴 `.to_dict()` 之外的屬性存取），與 `fill.py` 相同的相依方向。
"""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

SHEET_TITLE = "計算依據"

#: 欄位定義：(標題, 欄寬)
COLUMNS: tuple[tuple[str, int], ...] = (
    ("比較標的", 11),
    ("主要項目", 16),
    ("修正細項", 34),
    ("基準表頁", 9),
    ("比準地量測值", 16),
    ("比準地等級", 11),
    ("標的量測值", 16),
    ("標的等級", 10),
    ("修正百分比", 11),
    ("計入小計", 9),
    ("判級依據（比準地）", 34),
    ("判級依據（比較標的）", 34),
    ("矩陣查表", 30),
    ("敘述", 96),
)

_HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
_GROUP_FILL = PatternFill("solid", fgColor="F2F2F2")
_TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")


def add_evidence_sheet(
    wb: Workbook,
    evidence: list[Any],
    *,
    case_id: str,
    ruleset_id: str,
    notes: list[str] | None = None,
) -> Worksheet:
    """新增（或覆寫）「計算依據」工作表。

    `evidence` 是 `kernel.build_evidence()` 的回傳值。
    """
    if SHEET_TITLE in wb.sheetnames:
        del wb[SHEET_TITLE]
    ws = wb.create_sheet(SHEET_TITLE)

    row = _write_head(ws, case_id=case_id, ruleset_id=ruleset_id, notes=notes or [])
    header_row = row
    for idx, (title, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=idx, value=title)
        c.font = Font(bold=True)
        c.fill = _HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    row = header_row + 1
    for seg_ev in evidence:
        for fe in seg_ev.factors:
            _write_factor(ws, row, seg_ev.segment, fe)
            row += 1
        for g in seg_ev.groups:
            _write_group(ws, row, g)
            row += 1
        _write_total(ws, row, seg_ev)
        row += 2

    return ws


def _write_head(ws: Worksheet, *, case_id: str, ruleset_id: str, notes: list[str]) -> int:
    ws["A1"] = "計算依據明細"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"案號：{case_id}"
    ws["A3"] = f"規則集：{ruleset_id}"
    ws["A4"] = (
        "本表的每一列都由規則引擎產生，敘述不經過任何模型。"
        "等級依評價基準明細表的級距條文判定，修正百分比查該表的修正矩陣，"
        "頁碼欄可對回基準表原文。"
    )
    ws["A4"].alignment = Alignment(wrap_text=True, vertical="top")
    row = 5
    for n in notes:
        ws.cell(row=row, column=1, value=n).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    return row + 1


def _write_factor(ws: Worksheet, row: int, segment: str, fe: Any) -> None:
    values = (
        segment,
        fe.group,
        fe.label,
        fe.source_page,
        _display(fe.benchmark_value),
        _grade_text(fe.benchmark_grade, fe.benchmark_label),
        _display(fe.comparable_value),
        _grade_text(fe.comparable_grade, fe.comparable_label),
        _pct(fe.correction_pct),
        "是" if fe.counted else f"否（{fe.exclusion_reason}）",
        fe.benchmark_reason,
        fe.comparable_reason,
        fe.correction_reason,
        fe.narrative(),
    )
    for idx, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=idx, value=v)
        c.alignment = Alignment(vertical="top", wrap_text=idx in (3, 11, 12, 13, 14))


def _write_group(ws: Worksheet, row: int, g: Any) -> None:
    ws.cell(row=row, column=2, value=g.group).font = Font(bold=True)
    ws.cell(row=row, column=3, value="百分比小計").font = Font(bold=True)
    ws.cell(row=row, column=9, value=_pct(g.subtotal_pct)).font = Font(bold=True)
    c = ws.cell(row=row, column=14, value=g.narrative())
    c.alignment = Alignment(vertical="top", wrap_text=True)
    for idx in range(1, len(COLUMNS) + 1):
        ws.cell(row=row, column=idx).fill = _GROUP_FILL


def _write_total(ws: Worksheet, row: int, seg_ev: Any) -> None:
    ws.cell(row=row, column=1, value=seg_ev.segment).font = Font(bold=True)
    ws.cell(row=row, column=3, value="影響地價區域因素總修正數").font = Font(bold=True)
    ws.cell(row=row, column=9, value=_pct(seg_ev.total_pct)).font = Font(bold=True)
    c = ws.cell(row=row, column=14, value=seg_ev.narrative())
    c.alignment = Alignment(vertical="top", wrap_text=True)
    for idx in range(1, len(COLUMNS) + 1):
        ws.cell(row=row, column=idx).fill = _TOTAL_FILL


def _grade_text(grade: int | None, label: str) -> str:
    """等級欄。不適用者填「-」，與書表上的填法一致。"""
    return "-" if grade is None else f"{grade}（{label}）"


def _display(value: Any) -> str:
    return "無" if value is None else str(value)


def _pct(pct: Any) -> str:
    """百分比存成文字，因為這張表是給人讀的，不參與運算。

    零不帶正號，與敘述一致。
    """
    from decimal import Decimal

    q = Decimal(str(pct)).quantize(Decimal("0.01"))
    return f"{q}%" if q == 0 else f"{q:+}%"
