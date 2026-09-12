"""「計算依據」工作表的測試。

官方書表只有數字，沒有地方寫「這一格為什麼是這個值」。這張表把依據攤出來並與
書表放在同一個檔案，審查或訴願時不會分家。

範本不在版控（根目錄 .gitignore 排除 *.xlsx），找不到時整份 skip。
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import openpyxl
import pytest

from xlsxform import layout
from xlsxform.cli import compute_all, find_template, write_table5
from xlsxform.evidence_sheet import COLUMNS, SHEET_TITLE

ROOT = Path(__file__).resolve().parent.parent.parent
FACTS = json.loads(
    (ROOT / "kernel" / "fixtures" / "shulin_survey_facts.json").read_text(encoding="utf-8")
)
EXPECTED = FACTS["expected"]
COMPS = list(FACTS["comparables"])


def _template_dir() -> Path | None:
    env = os.environ.get("SHULIN_TEMPLATE_DIR")
    for c in ([Path(env)] if env else []) + [ROOT.parent / "正式題目"]:
        try:
            find_template(c, "table5")
        except (FileNotFoundError, OSError):
            continue
        return c
    return None


TEMPLATES = _template_dir()
pytestmark = pytest.mark.skipif(TEMPLATES is None, reason="找不到官方 xlsx 範本（不在版控裡）")


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    out = tmp_path_factory.mktemp("ev")
    computed = compute_all(FACTS)
    path, _ = write_table5(
        FACTS, computed, find_template(TEMPLATES, "table5"), out / "t5.xlsx", live=False
    )
    wb = openpyxl.load_workbook(path)
    return wb[SHEET_TITLE]


def _header_row(ws) -> int:
    for r in range(1, 20):
        if ws.cell(row=r, column=1).value == COLUMNS[0][0]:
            return r
    raise AssertionError("找不到表頭列")


def _rows(ws) -> list[dict]:
    hdr = _header_row(ws)
    titles = [t for t, _ in COLUMNS]
    out = []
    for r in range(hdr + 1, ws.max_row + 1):
        values = {t: ws.cell(row=r, column=i).value for i, t in enumerate(titles, start=1)}
        if any(v is not None for v in values.values()):
            out.append(values)
    return out


# ---------- 存在與結構 ----------


def test_sheet_is_added_next_to_the_form(sheet):
    wb = sheet.parent
    visible = [w.title for w in wb.worksheets if w.sheet_state != "hidden"]
    assert layout.SHEET_TABLE5_1 in visible
    assert SHEET_TITLE in visible


def test_official_form_sheet_is_untouched(sheet):
    """加依據表不能動到官方書表那一張。"""
    ws = sheet.parent[layout.SHEET_TABLE5_1]
    assert ws[layout.TABLE5_1_CASE_ID].value == FACTS["case_id"]
    assert ws["B42"].value == "=(1)+(2)+(3)+(4)+(5)+(6)+(7)+(8)"


def test_hidden_template_sheets_are_untouched(sheet):
    hidden = [w for w in sheet.parent.worksheets if w.sheet_state == "hidden"]
    assert len(hidden) == 23


def test_header_declares_no_model_involved(sheet):
    """要寫明敘述不經過模型。這是可舉證性的前提，也是紅線。"""
    head = "".join(str(sheet.cell(row=r, column=1).value or "") for r in range(1, 8))
    assert "不經過任何模型" in head
    assert FACTS["case_id"] in head
    assert "shulin-residential-regional" in head


def test_all_columns_present(sheet):
    hdr = _header_row(sheet)
    for i, (title, _) in enumerate(COLUMNS, start=1):
        assert sheet.cell(row=hdr, column=i).value == title


# ---------- 列數 ----------


def test_row_counts(sheet):
    """87 個細項列（29 × 3 標的）、24 個小計列（8 群組 × 3）、3 個總修正數列。"""
    rows = _rows(sheet)
    factor_rows = [r for r in rows if r["修正細項"] not in (None, "百分比小計", "影響地價區域因素總修正數")]
    subtotal_rows = [r for r in rows if r["修正細項"] == "百分比小計"]
    total_rows = [r for r in rows if r["修正細項"] == "影響地價區域因素總修正數"]
    assert len(factor_rows) == 87
    assert len(subtotal_rows) == 24
    assert len(total_rows) == 3


def test_every_comparable_appears(sheet):
    rows = _rows(sheet)
    segs = {r["比較標的"] for r in rows if r["比較標的"]}
    assert segs == set(COMPS)


# ---------- 內容 ----------


def test_every_factor_row_has_narrative_and_page(sheet):
    rows = [
        r
        for r in _rows(sheet)
        if r["修正細項"] not in (None, "百分比小計", "影響地價區域因素總修正數")
    ]
    missing_narrative = [r["修正細項"] for r in rows if not r["敘述"]]
    assert missing_narrative == []
    missing_page = [r["修正細項"] for r in rows if r["基準表頁"] is None]
    assert missing_page == []


def test_main_road_row_shows_the_whole_chain(sheet):
    """一列要交代完整的鏈：量測值、等級、判級依據、矩陣查表、修正率。"""
    row = next(
        r
        for r in _rows(sheet)
        if r["修正細項"] == "主要道路寬度" and r["比較標的"] == "P002-00"
    )
    assert row["比準地量測值"] == "28"
    assert row["比準地等級"] == "1（優）"
    assert row["標的量測值"] == "7"
    assert row["標的等級"] == "5（劣）"
    assert row["修正百分比"] == "+15.00%"
    assert row["基準表頁"] == 2
    assert "28m以上" in row["判級依據（比準地）"]
    assert "未滿8m" in row["判級依據（比較標的）"]
    assert "矩陣" in row["矩陣查表"]
    assert row["計入小計"] == "是"


def test_excluded_factor_says_why_not_counted(sheet):
    row = next(
        r for r in _rows(sheet) if r["修正細項"] == "容積率" and r["比較標的"] == "P002-00"
    )
    assert row["計入小計"].startswith("否")
    assert "併同" in row["計入小計"]
    assert "原載 260%" in row["敘述"], "追溯需要這一句"


def test_not_applicable_factor_shows_dash(sheet):
    row = next(
        r
        for r in _rows(sheet)
        if r["修正細項"] == "其他影響因素" and r["比較標的"] == "P002-00"
    )
    assert row["標的等級"] == "-"
    assert row["計入小計"].startswith("否")


def test_absent_factor_shows_none_not_blank(sheet):
    """未勾選要顯示「無」。空白會讓人以為漏填。"""
    row = next(
        r
        for r in _rows(sheet)
        if r["修正細項"].startswith("接近學校") and r["比較標的"] == "P002-00"
    )
    assert row["標的量測值"] == "無"
    assert row["標的等級"] == "5（劣）"


@pytest.mark.parametrize("seg", COMPS)
def test_total_row_matches_verified_value(seg, sheet):
    row = next(
        r
        for r in _rows(sheet)
        if r["修正細項"] == "影響地價區域因素總修正數" and r["比較標的"] == seg
    )
    expect = EXPECTED["total_correction_pct"][seg]
    assert row["修正百分比"] == f"+{expect}%"
    assert "交通運輸(2)" in row["敘述"]


def test_transport_subtotal_lists_its_members(sheet):
    row = next(
        r
        for r in _rows(sheet)
        if r["修正細項"] == "百分比小計" and r["主要項目"] == "交通運輸(2)"
    )
    assert row["修正百分比"] == "+23.50%"
    assert "6 個細項相加" in row["敘述"]


def test_zero_percent_has_no_plus_sign(sheet):
    """零不帶正號，與敘述一致。87 格裡多數是 0，加號會變雜訊。"""
    rows = _rows(sheet)
    zeros = [r["修正百分比"] for r in rows if r["修正百分比"] == "0.00%"]
    assert zeros, "應該有為 0 的列"
    assert not [r for r in rows if r["修正百分比"] == "+0.00%"]
