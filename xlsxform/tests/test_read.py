"""讀取填好的表3 xlsx，以及 round-trip 驗收。

round-trip 是這份測試的核心：把 fixtures 的事實寫成 xlsx，再讀回來，
除了 case_overrides 涵蓋的細項之外必須完全一致。它把「寫入」與「讀取」綁在
同一份格位對映上，任一邊改壞另一邊就會發現。

官方 xlsx 範本被根目錄 .gitignore 的 *.xlsx 排除（不在版控裡），
所以找不到範本時整份 skip。用環境變數 SHULIN_TEMPLATE_DIR 指定，
預設找 workspace 外層的「正式題目」。
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import pytest

from xlsxform import layout
from xlsxform.cli import compute_all, find_template, load_facts, write_table3
from xlsxform.read import SurveyReadError, apply_case_overrides, read_table3

ROOT = Path(__file__).resolve().parent.parent.parent
FACTS_PATH = ROOT / "kernel" / "fixtures" / "shulin_survey_facts.json"
FACTS = json.loads(FACTS_PATH.read_text(encoding="utf-8"))
EXPECTED = FACTS["expected"]
SEGMENTS = [FACTS["benchmark"]] + list(FACTS["comparables"])
OVERRIDES = FACTS["case_overrides"]
OVERRIDDEN_IDS = {o["factor_id"] for o in OVERRIDES}


def _template_dir() -> Path | None:
    env = os.environ.get("SHULIN_TEMPLATE_DIR")
    for c in ([Path(env)] if env else []) + [ROOT.parent / "正式題目"]:
        try:
            find_template(c, "table3")
        except (FileNotFoundError, OSError):
            continue
        return c
    return None


TEMPLATES = _template_dir()
pytestmark = pytest.mark.skipif(
    TEMPLATES is None,
    reason="找不到官方 xlsx 範本（不在版控裡）。設 SHULIN_TEMPLATE_DIR 指向範本目錄。",
)


@pytest.fixture(scope="module")
def written(tmp_path_factory) -> Path:
    """用 fixtures 的事實產出一份填好的表3 xlsx。"""
    out = tmp_path_factory.mktemp("t3")
    return write_table3(FACTS, find_template(TEMPLATES, "table3"), out / "t3.xlsx")


@pytest.fixture(scope="module")
def read_back(written):
    r = read_table3(written)
    return apply_case_overrides(r["segments"], OVERRIDES), r["warnings"]


# ---------- 基本讀取 ----------


def test_reads_all_four_segments(read_back):
    segments, _ = read_back
    assert sorted(segments) == sorted(SEGMENTS)


def test_no_warnings_for_our_own_output(read_back):
    """自己產出的檔案讀回來不該有任何警告。有的話是寫入或讀取有缺。"""
    _, warnings = read_back
    assert warnings == []


def test_source_is_recorded(read_back):
    segments, _ = read_back
    for seg in SEGMENTS:
        src = segments[seg]["source"]
        assert src["file"].endswith(".xlsx")
        assert seg in src["sheet"]


# ---------- round-trip ----------


@pytest.mark.parametrize("seg", SEGMENTS)
def test_roundtrip_facts_match_fixtures(seg, read_back):
    """套用 case_overrides 後的 facts 必須與 fixtures 完全一致，29 項全部。"""
    segments, _ = read_back
    expect = FACTS["segments"][seg]["facts"]
    got = segments[seg]["facts"]
    diff = {k: (got.get(k), v) for k, v in expect.items() if got.get(k) != v}
    assert diff == {}, f"{seg} 有 {len(diff)} 項不符（讀到, 期望）：{diff}"
    assert set(got) == set(expect), "細項集合也要一致"


@pytest.mark.parametrize("seg", SEGMENTS)
def test_roundtrip_raw_keeps_original_value(seg, read_back):
    """raw 要保留勘查表原載值，不可被覆寫污染。

    容積率是本案唯一有落差的細項：勘查表原載 260%（P002-00 是 200%），
    計算依局處指示一律用 200%。
    """
    segments, _ = read_back
    raw = segments[seg]["raw"]
    fid = "regional.land_control.floor_area_ratio"
    assert raw[fid] == FACTS["segments"][seg]["raw"][fid]
    assert segments[seg]["facts"][fid] == 200


@pytest.mark.parametrize("seg", SEGMENTS)
def test_roundtrip_extras(seg, read_back):
    """路名與土地改良勾選項目要讀得回來。

    這兩樣不是評價細項，所以不在 raw 裡。漏了它們 round-trip 會掉字：
    reader 產出的 raw 是型別轉換後的數值，解析不出路名（實測踩過，
    第二次產出時 G11 與 E31／E32 會變空白）。
    """
    segments, _ = read_back
    extras = segments[seg]["extras"]
    assert extras["main_road_name"], f"{seg} 讀不到主要道路名稱"
    assert len(extras["improvement_items"]) == FACTS["segments"][seg]["facts"][
        "regional.land_improvement.site_improvement"
    ]


def test_roundtrip_segment_range(read_back):
    segments, _ = read_back
    for seg in SEGMENTS:
        assert segments[seg]["segment_range"] == FACTS["segments"][seg]["segment_range"]


# ---------- 兩條輸入路徑產生相同結果 ----------


def test_both_input_paths_produce_same_numbers(written):
    """--facts 與 --from-xlsx 跑同一條鏈，數字必須相同。"""
    from_json = compute_all(FACTS)
    from_xlsx = compute_all(load_facts(FACTS_PATH, written))

    for seg in FACTS["comparables"]:
        assert from_xlsx["table5_1"].totals[seg] == from_json["table5_1"].totals[seg]
        for key in ("regional_pct", "abs_sum_pct", "trial_price", "weight_pct", "similarity"):
            assert from_xlsx["table4"][key][seg] == from_json["table4"][key][seg], key

    for key in ("benchmark_comparison_price", "benchmark_land_price"):
        assert from_xlsx["table4"][key] == from_json["table4"][key]

    # 順便確認還是那組已驗證的數字，不是兩邊一起錯
    assert from_xlsx["table4"]["benchmark_comparison_price"] == EXPECTED[
        "benchmark_comparison_price"
    ]
    assert from_xlsx["table4"]["benchmark_land_price"] == EXPECTED["benchmark_land_price"]


def test_load_facts_rejects_mismatched_segments(written, tmp_path):
    """上傳的 xlsx 若不是本案的勘查表，要明確拒絕而不是算下去。"""
    settings = json.loads(FACTS_PATH.read_text(encoding="utf-8"))
    settings["comparables"] = ["P099-00"]
    p = tmp_path / "settings.json"
    p.write_text(json.dumps(settings, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(SystemExit, match="不符"):
        load_facts(p, written)


# ---------- case_overrides ----------


def test_override_applies_to_all_segments(read_back):
    segments, _ = read_back
    for seg in SEGMENTS:
        for fid in OVERRIDDEN_IDS:
            assert segments[seg]["facts"][fid] == 200


def test_override_rejects_unknown_segment(read_back):
    segments, _ = read_back
    with pytest.raises(SurveyReadError, match="不存在的區段"):
        apply_case_overrides(
            segments,
            [{"factor_id": "regional.land_control.floor_area_ratio", "applies_to": ["P099-00"], "value": 1}],
        )


def test_override_has_source_and_reason():
    """覆寫必須說明來源與理由。徵收案會被訴願，「為什麼用 200 不用 260」要答得出來。"""
    for ov in OVERRIDES:
        assert ov.get("source"), ov
        assert ov.get("reason"), ov
        assert "factor_id" in ov and "value" in ov


# ---------- 型別轉換 ----------


def test_percent_from_string_and_formatted_number():
    from xlsxform.read import _to_percent

    assert _to_percent("50%", None, where="t") == 50
    assert _to_percent("50％", None, where="t") == 50  # 全角
    assert _to_percent(0.5, "0.00%", where="t") == 50
    assert _to_percent(50, "General", where="t") == 50
    assert _to_percent(None, None, where="t") is None


def test_percent_refuses_to_guess_ambiguous_value():
    """數值小於 1 而格式不含 % 時無法判斷是 0.5 個百分點還是 50%，要報錯不猜。"""
    from xlsxform.read import _to_percent

    with pytest.raises(SurveyReadError, match="無法判斷"):
        _to_percent(0.5, "General", where="P001-00 H7")


def test_number_strips_unit():
    from xlsxform.read import _to_number

    assert _to_number(28) == 28
    assert _to_number("28") == 28
    assert _to_number("28M") == 28
    assert _to_number("1,234") == 1234
    assert _to_number(7.5) == 7.5
    assert _to_number(None) is None


def test_absent_tokens_do_not_include_none_of_the_meaningful_ones():
    """「無」不可視為空白。

    它在有無禁止建築、有無限制建築是有效答案（規則集對映到第 1 級）。
    正規化成 None 會讓那兩格讀不到值，實測踩過。
    """
    from xlsxform.read import _ABSENT_TOKENS, _to_text

    assert "無" not in _ABSENT_TOKENS
    assert _to_text("無") == "無"
    assert _to_text("○") is None
    assert _to_text("") is None
    assert _to_text("  ") is None


# ---------- 工作表辨識 ----------


def test_segment_no_from_sheet_title(written):
    """xlsxform 產出的工作表叫「表3 P001-00」，從名稱抓區段編號。"""
    import openpyxl

    from xlsxform.read import _segment_no

    wb = openpyxl.load_workbook(written)
    ws = wb[layout.TABLE3_SHEET_TITLE.format(segment="P003-00")]
    assert _segment_no(ws) == "P003-00"


def test_segment_no_falls_back_to_cell(tmp_path):
    """官方空白範本只有一張叫「表3區段勘查表」，要能從 G3 讀出編號。"""
    import openpyxl

    from xlsxform.read import _segment_no

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = layout.SHEET_TABLE3
    ws[layout.TABLE3_CELL_SEGMENT_NO] = "P007-00"
    assert _segment_no(ws) == "P007-00"


def test_segment_no_raises_when_unavailable():
    import openpyxl

    from xlsxform.read import _segment_no

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "沒有編號的表"
    with pytest.raises(SurveyReadError, match="找不到地價區段編號"):
        _segment_no(ws)


def test_missing_file_raises():
    with pytest.raises(SurveyReadError, match="找不到檔案"):
        read_table3("/tmp/不存在的檔案-xyz.xlsx")
