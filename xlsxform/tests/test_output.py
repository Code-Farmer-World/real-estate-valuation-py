"""產出的 xlsx 內容驗收。

最重要的一條是活版與定版的一致性：活版把加總與四則運算寫成 Excel 公式，
定版寫規則引擎算好的數值。兩者算出來必須相同，否則有一邊寫錯了。

openpyxl 寫入公式後檔案裡沒有 cached value，所以無法直接讀公式的計算結果。
這裡的做法是把公式描述的關係套用在定版的數值上驗證，例如活版的
G18 是 =SUM(G12:G17)，就檢查定版的 G18 是否等於定版 G12 到 G17 的和。
Excel 開啟活版時會得到同樣的結果。

官方 xlsx 範本被根目錄 .gitignore 的 *.xlsx 排除（不在版控裡），
所以找不到範本時整份 skip，與 parser 測試依賴 VALUATION_DOC_DIR 的處理一致。
用環境變數 SHULIN_TEMPLATE_DIR 指定，預設找 workspace 外層的「正式題目」。
"""

from __future__ import annotations

import json
import os
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from xlsxform import layout
from xlsxform.cli import compute_all, find_template, write_table3, write_table4, write_table5
from xlsxform.fill import checked_improvements

ROOT = Path(__file__).resolve().parent.parent.parent
FACTS_PATH = ROOT / "kernel" / "fixtures" / "shulin_survey_facts.json"
FACTS = json.loads(FACTS_PATH.read_text(encoding="utf-8"))
EXPECTED = FACTS["expected"]
COMPS = list(FACTS["comparables"])


def _template_dir() -> Path | None:
    env = os.environ.get("SHULIN_TEMPLATE_DIR")
    candidates = [Path(env)] if env else []
    candidates.append(ROOT.parent / "正式題目")
    for c in candidates:
        try:
            find_template(c, "table5")
        except (FileNotFoundError, OSError):
            continue
        return c
    return None


TEMPLATES = _template_dir()
pytestmark = pytest.mark.skipif(
    TEMPLATES is None,
    reason="找不到官方 xlsx 範本（不在版控裡）。設 SHULIN_TEMPLATE_DIR 指向範本目錄。",
)


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("forms")
    computed = compute_all(FACTS)
    paths = {
        "table3": write_table3(FACTS, find_template(TEMPLATES, "table3"), out / "t3.xlsx"),
    }
    for live, tag in ((True, "live"), (False, "final")):
        p, counts = write_table5(
            FACTS, computed, find_template(TEMPLATES, "table5"), out / f"t5-{tag}.xlsx", live=live
        )
        paths[f"table5_{tag}"] = p
        paths[f"table5_{tag}_counts"] = counts
        paths[f"table4_{tag}"] = write_table4(
            computed, find_template(TEMPLATES, "table4"), out / f"t4-{tag}.xlsx", live=live
        )
    return paths


def _sheet(path, title=None):
    wb = openpyxl.load_workbook(path)
    vis = [w for w in wb.worksheets if w.sheet_state != "hidden"]
    if title is None:
        assert len(vis) == 1, [w.title for w in vis]
        return vis[0]
    return wb[title]


def _v(ws, coord):
    return ws[coord].value


def _d(ws, coord) -> Decimal:
    """讀數值並轉 Decimal。百分比欄位存的是小數（0.235 代表 23.5%）。"""
    raw = ws[coord].value
    return Decimal(str(raw if raw is not None else 0))


# ---------- layout ----------


def test_layout_self_check():
    assert layout.check_layout() == []


def test_layout_factor_ids_match_ruleset():
    from kernel.src.ruleset import load_ruleset

    rs = load_ruleset(FACTS["ruleset_regional"])
    rows = layout.TABLE5_1_FACTOR_ROWS
    assert [rows[r] for r in sorted(rows)] == rs.factor_ids


# ---------- 表5-1 ----------


def test_table5_1_cell_counts(built):
    """優劣等級是兩欄：級數與等級文字。

    新北查估書表製作手冊第 5 章第 42 頁：「左欄填載各該地價區段之優劣等級
    級數，右欄填載優劣等級細項」。表頭 C4:D4 是合併格所以看起來像一欄，
    細項列的 C 與 D 各自獨立。先前只填級數，漏了 116 格等級文字。
    """
    assert built["table5_final_counts"] == {
        "grades": 116,
        "grade_labels": 116,
        "corrections": 87,
        "subtotals": 24,
        "totals": 3,
    }


def test_table5_1_grade_labels_pair_with_the_grade_numbers(built):
    """每一格級數旁邊都要有對應的等級文字，而且文字要與規則集的語彙一致。

    官方已填好的金山範本逐列都是「1 優」「3 普通」「5 劣」這種兩欄形式，
    題目的表5-1 在其他影響因素那列自己填的是「- 無」。
    """
    ws = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    pairs = [
        (layout.TABLE5_1_GRADE_COL["benchmark"], layout.TABLE5_1_GRADE_LABEL_COL["benchmark"])
    ] + [
        (layout.TABLE5_1_GRADE_COL[i], layout.TABLE5_1_GRADE_LABEL_COL[i])
        for i in range(len(COMPS))
    ]

    seen = 0
    for row in layout.TABLE5_1_FACTOR_ROWS:
        for grade_col, label_col in pairs:
            grade = _v(ws, f"{grade_col}{row}")
            label = _v(ws, f"{label_col}{row}")
            assert grade is not None, f"{grade_col}{row} 級數空白"
            assert label, f"{label_col}{row} 等級文字空白"
            if grade == "-":
                # 不適用。題目自己填的是「- 無」
                assert label == "無", f"{label_col}{row} 應為「無」，實得 {label!r}"
            else:
                assert label != "-", f"{label_col}{row} 不該是「-」"
            seen += 1
    assert seen == 116


def test_table5_1_header(built):
    ws = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    assert _v(ws, layout.TABLE5_1_CASE_ID) == FACTS["case_id"]
    row = layout.TABLE5_1_SEGMENT_ROW
    assert _v(ws, f"C{row}") == FACTS["benchmark"]
    for i, seg in enumerate(COMPS):
        assert _v(ws, f"{layout.TABLE5_1_GRADE_COL[i]}{row}") == seg


def test_table5_1_every_grade_cell_filled(built):
    ws = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    cols = ["C"] + [layout.TABLE5_1_GRADE_COL[i] for i in range(len(COMPS))]
    blanks = [
        f"{c}{row}"
        for row in layout.TABLE5_1_FACTOR_ROWS
        for c in cols
        if _v(ws, f"{c}{row}") in (None, "")
    ]
    assert blanks == [], f"這些等級欄是空的：{blanks}"


def test_table5_1_not_applicable_is_dash(built):
    ws = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    rows = {v: k for k, v in layout.TABLE5_1_FACTOR_ROWS.items()}
    for fid in FACTS["not_applicable_factors"]:
        row = rows[fid]
        for c in ["C"] + [layout.TABLE5_1_GRADE_COL[i] for i in range(len(COMPS))]:
            assert _v(ws, f"{c}{row}") == "-", f"{c}{row}"


@pytest.mark.parametrize("seg_index,seg", list(enumerate(COMPS)))
def test_table5_1_total_matches(built, seg_index, seg):
    ws = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    col = layout.TABLE5_1_PCT_COL[seg_index]
    got = _d(ws, f"{col}{layout.TABLE5_1_TOTAL_ROW}") * 100
    assert got == Decimal(EXPECTED["total_correction_pct"][seg])


def test_table5_1_b42_untouched(built):
    """B42 印的 =(1)+(2)+...+(8) 是說明文字不是公式，不可覆寫。"""
    ws = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    assert _v(ws, "B42") == "=(1)+(2)+(3)+(4)+(5)+(6)+(7)+(8)"


def test_table5_1_live_has_formulas(built):
    ws = _sheet(built["table5_live"], layout.SHEET_TABLE5_1)
    for row in layout.TABLE5_1_SUBTOTAL_ROW_ORDER:
        assert str(_v(ws, f"G{row}")).startswith("=SUM("), row
    assert str(_v(ws, f"G{layout.TABLE5_1_TOTAL_ROW}")).startswith("=SUM(")
    # 等級與修正率在活版仍是數值，那些是查表判斷的結果不是算術
    assert _v(ws, "C12") == "1"
    assert isinstance(_v(ws, "G12"), (int, float))


def test_table5_1_land_control_subtotal_skips_three_rows(built):
    """第 11 列的公式必須跳過 6、7、8（使用分區、建蔽率、容積率）。"""
    ws = _sheet(built["table5_live"], layout.SHEET_TABLE5_1)
    assert _v(ws, "G11") == "=SUM(G5,G9,G10)"


# ---------- 活版公式與定版數值的一致性 ----------


def test_subtotal_formula_matches_final_values(built):
    """活版 =SUM(範圍) 套在定版數值上，要等於定版的小計。"""
    final = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    for i in range(len(COMPS)):
        col = layout.TABLE5_1_PCT_COL[i]
        for row, member_rows in layout.TABLE5_1_SUBTOTAL_ROWS.items():
            expect = sum((_d(final, f"{col}{r}") for r in member_rows), Decimal(0))
            assert _d(final, f"{col}{row}") == expect, f"{col}{row}"


def test_total_formula_matches_final_values(built):
    final = _sheet(built["table5_final"], layout.SHEET_TABLE5_1)
    for i in range(len(COMPS)):
        col = layout.TABLE5_1_PCT_COL[i]
        expect = sum(
            (_d(final, f"{col}{r}") for r in layout.TABLE5_1_SUBTOTAL_ROW_ORDER), Decimal(0)
        )
        assert _d(final, f"{col}{layout.TABLE5_1_TOTAL_ROW}") == expect


def test_table4_abs_sum_formula_matches_final_values(built):
    """活版 =ABS(J6)+ABS(J8)+SUMPRODUCT(ABS(J9:J28))，個別因素為空時就是前兩項。"""
    final = _sheet(built["table4_final"], layout.SHEET_TABLE4)
    for i in range(len(COMPS)):
        diff = layout.TABLE4_DIFF_COL[i]
        left = layout.TABLE4_DECISION_LEFT_COL[i]
        expect = abs(_d(final, f"{diff}{layout.TABLE4_ROW_TRANSACTION_DATE}")) + abs(
            _d(final, f"{diff}{layout.TABLE4_ROW_SEGMENT}")
        )
        assert _d(final, f"{left}{layout.TABLE4_ROW_ABS_SUM}") == expect


def test_table4_trial_price_formula_matches_final_values(built):
    """活版 =ROUND(G7*(1+J8)*(1+G29),0)。"""
    final = _sheet(built["table4_final"], layout.SHEET_TABLE4)
    for i in range(len(COMPS)):
        cond = layout.TABLE4_COND_COL[i]
        diff = layout.TABLE4_DIFF_COL[i]
        left = layout.TABLE4_DECISION_LEFT_COL[i]
        price = _d(final, f"{cond}{layout.TABLE4_ROW_ADJUSTED_UNIT_PRICE}")
        regional = _d(final, f"{diff}{layout.TABLE4_ROW_SEGMENT}")
        individual = _d(final, f"{cond}{layout.TABLE4_ROW_INDIVIDUAL_TOTAL}")
        expect = (price * (1 + regional) * (1 + individual)).quantize(Decimal(1))
        assert _d(final, f"{left}{layout.TABLE4_ROW_TRIAL_PRICE}") == expect


def test_table4_benchmark_price_formula_matches_final_values(built):
    """活版 =ROUND(G31*I31+K31*M31+O31*Q31,0)。"""
    final = _sheet(built["table4_final"], layout.SHEET_TABLE4)
    trial_row = layout.TABLE4_ROW_TRIAL_PRICE
    acc = Decimal(0)
    for i in range(len(COMPS)):
        left = layout.TABLE4_DECISION_LEFT_COL[i]
        right = layout.TABLE4_DECISION_RIGHT_COL[i]
        acc += _d(final, f"{left}{trial_row}") * _d(final, f"{right}{trial_row}")
    got = _d(final, layout.TABLE4_BENCHMARK_PRICE_CELL)
    assert got == acc.quantize(Decimal(1))
    assert int(got) == EXPECTED["benchmark_comparison_price"]


# ---------- 表4 ----------


def test_table4_given_values_are_copied(built):
    final = _sheet(built["table4_final"], layout.SHEET_TABLE4)
    given = FACTS["table4_given"]["segments"]
    for i, seg in enumerate(COMPS):
        cond = layout.TABLE4_COND_COL[i]
        assert _v(final, f"{cond}{layout.TABLE4_ROW_NORMAL_UNIT_PRICE}") == given[seg][
            "normal_unit_price"
        ]
        assert _v(final, f"{cond}{layout.TABLE4_ROW_ADJUSTED_UNIT_PRICE}") == given[seg][
            "adjusted_unit_price"
        ]
        assert _v(final, f"{cond}{layout.TABLE4_ROW_SEGMENT}") == seg


def test_table4_header_is_filled(built):
    """表頭六格：估價基準日、案號、比準地宗地流水號、三個實例編號。

    這六格題目都有給，先前漏填。逐格盤點產出檔案時才發現，所以留這個測試。
    一律是文字：流水號 0003 被當數字就會掉前導零。
    """
    g = FACTS["table4_given"]
    for final in (
        _sheet(built["table4_final"], layout.SHEET_TABLE4),
        _sheet(built["table4_live"], layout.SHEET_TABLE4),
    ):
        assert _v(final, layout.TABLE4_CELL_BASE_DATE) == str(g["appraisal_base_date"])
        assert _v(final, layout.TABLE4_CELL_CASE_ID) == str(g["case_id"])
        serial = _v(final, layout.TABLE4_CELL_BENCHMARK_SERIAL)
        assert serial == str(g["benchmark_parcel_serial"])
        assert serial.startswith("0"), "宗地流水號的前導零掉了，被當成數字了"

        for i, seg in enumerate(COMPS):
            assert _v(final, layout.TABLE4_EXAMPLE_NO_CELLS[i]) == str(
                g["segments"][seg]["example_no"]
            )


def test_table4_individual_factor_rows_left_blank(built):
    """項目 7 到 25 一律留空，題目未提供宗地個別條件資料。"""
    final = _sheet(built["table4_final"], layout.SHEET_TABLE4)
    for i in range(len(COMPS)):
        diff = layout.TABLE4_DIFF_COL[i]
        for row in layout.TABLE4_INDIVIDUAL_ROWS:
            assert _v(final, f"{diff}{row}") is None, f"{diff}{row} 應留空"


def test_table4_weights_and_similarity(built):
    final = _sheet(built["table4_final"], layout.SHEET_TABLE4)
    for i, seg in enumerate(COMPS):
        right = layout.TABLE4_DECISION_RIGHT_COL[i]
        exp = EXPECTED["weights"][seg]
        assert _v(final, f"{right}{layout.TABLE4_ROW_ABS_SUM}") == exp["similarity"]
        assert _d(final, f"{right}{layout.TABLE4_ROW_TRIAL_PRICE}") * 100 == Decimal(
            exp["weight_pct"]
        )


def test_table4_remark_states_individual_factor_premise(built):
    """全案備註必須寫明個別因素的前提，否則試算價格會被當成完整答案。"""
    final = _sheet(built["table4_final"], layout.SHEET_TABLE4)
    text = _v(final, layout.TABLE4_REMARK_ALL_CASE_CELL) or ""
    assert "個別因素" in text
    assert "地價查估單位" in text
    assert "0%" in text or "0％" in text
    assert f"{EXPECTED['benchmark_land_price']:,}" in text


# ---------- 表3 ----------


def test_table3_has_one_sheet_per_segment(built):
    wb = openpyxl.load_workbook(built["table3"])
    vis = [w.title for w in wb.worksheets if w.sheet_state != "hidden"]
    segments = [FACTS["benchmark"]] + COMPS
    assert vis == [layout.TABLE3_SHEET_TITLE.format(segment=s) for s in segments]
    hidden = [w for w in wb.worksheets if w.sheet_state == "hidden"]
    assert len(hidden) == 23, "隱藏的舊範本工作表不應被動到"


def test_table3_preserves_merged_cells(built):
    """複製工作表不能弄壞合併格。範本原本有 188 個。"""
    wb = openpyxl.load_workbook(built["table3"])
    for w in [x for x in wb.worksheets if x.sheet_state != "hidden"]:
        assert len(w.merged_cells.ranges) == 188, w.title


def test_table3_values_per_segment(built):
    wb = openpyxl.load_workbook(built["table3"])
    for seg in [FACTS["benchmark"]] + COMPS:
        ws = wb[layout.TABLE3_SHEET_TITLE.format(segment=seg)]
        d = FACTS["segments"][seg]
        assert _v(ws, layout.TABLE3_CELL_SEGMENT_NO) == seg
        assert _v(ws, layout.TABLE3_CELL_YEAR) == FACTS["year_period"]
        assert _v(ws, layout.TABLE3_CELL_SEGMENT_RANGE) == d["segment_range"]
        assert _v(ws, layout.TABLE3_CELL_MAIN_ROAD_WIDTH) == d["facts"][
            "regional.transport.main_road_width"
        ]
        assert _v(ws, layout.TABLE3_CELL_AVG_ROAD_WIDTH) == d["facts"][
            "regional.transport.avg_road_width"
        ]
        assert _v(ws, "I23") == d["facts"]["regional.transport.road_development"]


def test_table3_records_original_floor_area_ratio(built):
    """表3 填勘查表原載的容積率，不是計算用的 200%。

    容積率一律以 200% 計算是局處對「計算」的指示，勘查表本身要忠實記載現況。
    這個區分寫在表5-1 的備註欄。
    """
    wb = openpyxl.load_workbook(built["table3"])
    for seg in [FACTS["benchmark"]] + COMPS:
        ws = wb[layout.TABLE3_SHEET_TITLE.format(segment=seg)]
        original = FACTS["segments"][seg]["raw"]["regional.land_control.floor_area_ratio"]
        assert _v(ws, "H7") == f"{original}%", seg
        assert FACTS["segments"][seg]["facts"]["regional.land_control.floor_area_ratio"] == 200


def test_table3_keeps_printed_labels_and_writes_values_beside_them(built):
    """建築密度與建築型態的值要填在 R 欄，Q 欄那個標籤不能被蓋掉。

    先前誤把值寫進 Q42/Q43，結果「建築密度」「建築型態」兩個標籤被覆寫，
    表格看起來就少了欄位名。
    """
    wb = openpyxl.load_workbook(built["table3"])
    for seg in [FACTS["benchmark"]] + COMPS:
        ws = wb[layout.TABLE3_SHEET_TITLE.format(segment=seg)]
        only = FACTS["segments"][seg]["table3_only"]
        assert _v(ws, "Q42") == "建築密度", f"{seg} 的標籤被覆寫了"
        assert _v(ws, "Q43") == "建築型態", f"{seg} 的標籤被覆寫了"
        assert _v(ws, layout.TABLE3_CELL_BUILDING_DENSITY) == only["building_density"]
        assert _v(ws, layout.TABLE3_CELL_BUILDING_TYPE) == only["building_type"]


def test_table3_land_use_is_marked(built):
    """土地利用現況要把勾選的 ○ 改成 ●，未勾的維持 ○。

    P001-00 勾商業用與住宅用，其餘三段只勾住宅用。
    「●住宅用」不可以誤中「○住商混合」。
    """
    wb = openpyxl.load_workbook(built["table3"])
    for seg in [FACTS["benchmark"]] + COMPS:
        ws = wb[layout.TABLE3_SHEET_TITLE.format(segment=seg)]
        text = _v(ws, layout.TABLE3_CELL_LAND_USE) or ""
        expect = FACTS["segments"][seg]["table3_only"]["land_use_current"]
        assert text.count("●") == len(expect), seg
        for name in expect:
            assert f"●{name}" in text, f"{seg} 少勾 {name}"
        assert "○住商混合" in text, f"{seg} 誤勾了住商混合"


def test_table3_improvement_checkboxes_match_fact_count(built):
    """土地改良的勾選數必須等於事實記載的項數。

    這個項數直接決定該細項的優劣等級（四項以上為優），漏勾會讓等級判錯。
    先前的實作只依 ■ 切割字串，把「■開挖水溝 □水土保持」當成一個項目名，
    結果四項只勾到兩項。
    """
    wb = openpyxl.load_workbook(built["table3"])
    for seg in [FACTS["benchmark"]] + COMPS:
        ws = wb[layout.TABLE3_SHEET_TITLE.format(segment=seg)]
        text = (_v(ws, "E31") or "") + (_v(ws, "E32") or "")
        expect = FACTS["segments"][seg]["facts"]["regional.land_improvement.site_improvement"]
        assert text.count("■") == expect, seg


def test_checked_improvements_parses_mixed_marks():
    raw = "■整平或填挖基地 ■開挖水溝 □水土保持 ■鋪築道路 ■埋設管道 □修築駁嵌"
    assert checked_improvements(raw) == {
        "整平或填挖基地",
        "開挖水溝",
        "鋪築道路",
        "埋設管道",
    }


def test_checked_improvements_handles_empty():
    assert checked_improvements(None) == set()
    assert checked_improvements("") == set()
    assert checked_improvements("□全部未勾") == set()
