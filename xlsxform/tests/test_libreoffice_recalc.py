"""用 LibreOffice 的公式引擎重算活版，跟定版比對。

## 這組測試補上的是哪一個洞

`openpyxl` 寫入公式後檔案裡沒有快取值，程式讀回來只拿得到公式字串。
所以在這之前，「活版的 Excel 公式在真的試算表軟體裡會算出什麼」是完全沒驗過的。
`verify.py` 那 10 項驗的是「公式描述的關係」，做法是把關係用 Python 重算一次
套在定版數值上比對。那證明關係對，不證明 Excel 接受那些公式。

如果公式語法錯、引用的格位錯、或範圍裡混進文字讓 `SUMPRODUCT` 回 `#VALUE!`，
交出去的檔案打開會是一片錯誤訊息，而所有測試都還是綠的。

LibreOffice 轉存 xlsx 時會重算並寫入快取值，所以把活版丟給它跑一次，
再用 `data_only=True` 讀，就拿得到真實引擎算出來的數字。

找不到 LibreOffice 就整支跳過。安裝方式見 `xlsxform/topdf.py` 的模組說明。
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from xlsxform import layout
from xlsxform.pipeline import compute_all, find_template, write_table4, write_table5
from xlsxform.topdf import DEFAULT_EXCLUDE_SHEETS, find_soffice, recalculate, strip_hidden_sheets, to_pdf

ROOT = Path(__file__).resolve().parent.parent.parent
FACTS = json.loads((ROOT / "kernel" / "fixtures" / "shulin_survey_facts.json").read_text("utf-8"))
TEMPLATES = ROOT.parent / "正式題目"

SOFFICE = find_soffice()

pytestmark = [
    pytest.mark.skipif(not TEMPLATES.exists(), reason=f"找不到官方範本 {TEMPLATES}"),
    pytest.mark.skipif(
        SOFFICE is None, reason="找不到 LibreOffice，安裝方式見 xlsxform/topdf.py"
    ),
]

#: 表5-1 的小計列與總修正數列。這些格在活版是 `=SUM(...)`。
SUBTOTAL_ROWS = (11, 18, 24, 26, 33, 37, 39, 41)
TOTAL_ROW = 42
PERCENT_COLS = ("G", "J", "M")

#: 表4 的計算欄。活版分別是 SUM、ABS＋SUMPRODUCT、ROUND 那三種公式。
TABLE4_COMPUTED_CELLS = (
    "G29", "K29", "O29",
    "G30", "K30", "O30",
    "G31", "K31", "O31",
    "G32",
)


@pytest.fixture(scope="module")
def pair(tmp_path_factory):
    """產一組活版與定版，並讓 LibreOffice 重算活版。"""
    out = tmp_path_factory.mktemp("recalc")
    computed = compute_all(FACTS)

    made = {}
    for live, tag in ((True, "live"), (False, "final")):
        p, _ = write_table5(
            FACTS, computed, find_template(TEMPLATES, "table5"), out / f"t5-{tag}.xlsx", live=live
        )
        made[f"t5_{tag}"] = p
        made[f"t4_{tag}"] = write_table4(
            computed, find_template(TEMPLATES, "table4"), out / f"t4-{tag}.xlsx", live=live
        )

    done = out / "recalculated"
    made["t5_recalc"] = recalculate(made["t5_live"], done, soffice=SOFFICE)
    made["t4_recalc"] = recalculate(made["t4_live"], done, soffice=SOFFICE)
    return made


def _values(path: Path, title: str) -> dict:
    ws = openpyxl.load_workbook(path, data_only=True)[title]
    return ws


def test_table5_subtotals_recalculate_to_the_final_values(pair):
    """24 格群組小計與 3 格總修正數，真實引擎算出來的值要等於定版。"""
    recalc = _values(pair["t5_recalc"], layout.SHEET_TABLE5_1)
    final = _values(pair["t5_final"], layout.SHEET_TABLE5_1)

    checked = 0
    for row in list(SUBTOTAL_ROWS) + [TOTAL_ROW]:
        for col in PERCENT_COLS:
            coord = f"{col}{row}"
            got, want = recalc[coord].value, final[coord].value
            assert got is not None, f"{coord} 重算後是空的，公式可能沒被算"
            assert float(got) == pytest.approx(float(want)), coord
            checked += 1
    assert checked == 27


def test_table4_computed_cells_recalculate_to_the_final_values(pair):
    """表4 的十個計算欄，含 SUMPRODUCT 那三格。

    `SUMPRODUCT(ABS(J9:J28))` 是最脆弱的一個：那個範圍如果混進文字就會回
    `#VALUE!`，而範本在鄰近欄位印了單位「M」。本案個別因素全空，所以應該得 0。
    """
    recalc = _values(pair["t4_recalc"], layout.SHEET_TABLE4)
    final = _values(pair["t4_final"], layout.SHEET_TABLE4)

    for coord in TABLE4_COMPUTED_CELLS:
        got, want = recalc[coord].value, final[coord].value
        assert got is not None, f"{coord} 重算後是空的"
        assert not isinstance(got, str), f"{coord} 重算後是字串 {got!r}，公式可能出錯"
        assert float(got) == pytest.approx(float(want)), coord


def test_no_formula_evaluates_to_an_error(pair):
    """任何一格都不能算出 #VALUE! #REF! #NAME? 這類錯誤。

    LibreOffice 把錯誤寫成字串，所以掃一遍所有格找 `#` 開頭的字串就抓得到。
    """
    for key, title in (
        (pair["t5_recalc"], layout.SHEET_TABLE5_1),
        (pair["t4_recalc"], layout.SHEET_TABLE4),
    ):
        ws = openpyxl.load_workbook(key, data_only=True)[title]
        errors = [
            (c.coordinate, c.value)
            for row in ws.iter_rows()
            for c in row
            if isinstance(c.value, str) and c.value.startswith("#")
        ]
        assert not errors, f"{title} 有公式算出錯誤值：{errors}"


def test_strip_hidden_sheets_keeps_only_what_gets_printed(pair, tmp_path):
    """官方範本帶 23 張隱藏的舊範本樣例，LibreOffice 轉 PDF 時不理隱藏屬性。

    不先剝掉的話，表5-1 那份會從 1 頁變 58 頁，而且第 1 頁是別的表。
    """
    src = pair["t5_final"]
    before = openpyxl.load_workbook(src)
    assert sum(1 for w in before.worksheets if w.sheet_state == "hidden") == 23

    dest = tmp_path / "flat.xlsx"
    kept = strip_hidden_sheets(src, dest, exclude=DEFAULT_EXCLUDE_SHEETS)

    assert kept == [layout.SHEET_TABLE5_1]
    after = openpyxl.load_workbook(dest)
    assert after.sheetnames == [layout.SHEET_TABLE5_1]


def test_pdf_has_one_page_per_form(pair, tmp_path):
    """表5-1 與表4 各自一頁。頁數靠範本的列印版面設定，那些設定要沒被弄壞。"""
    pytest.importorskip("pdfplumber")
    import pdfplumber

    for key, expect_pages in ((pair["t5_final"], 1), (pair["t4_final"], 1)):
        pdf = to_pdf(key, tmp_path / "pdf", soffice=SOFFICE)
        with pdfplumber.open(pdf) as doc:
            assert len(doc.pages) == expect_pages, f"{pdf.name} 有 {len(doc.pages)} 頁"


def test_pdf_contains_the_key_numbers(pair, tmp_path):
    """轉出來的 PDF 要看得到算好的數字，不是一片空白或錯誤。"""
    pytest.importorskip("pdfplumber")
    import pdfplumber

    expected = FACTS["expected"]
    pdf = to_pdf(pair["t4_final"], tmp_path / "pdf", soffice=SOFFICE)
    with pdfplumber.open(pdf) as doc:
        text = doc.pages[0].extract_text() or ""

    assert FACTS["case_id"] in text
    assert f"{expected['benchmark_comparison_price']:,}" in text
    for seg in FACTS["comparables"]:
        assert f"{expected['trial_price'][seg]:,}" in text, seg
