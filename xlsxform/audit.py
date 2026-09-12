"""回填完整性盤點：把產出檔案跟原始範本逐格比，列出我們動了哪些格。

## 為什麼需要這個

這個專案的錯誤有兩種，性質完全不同：

**計算錯誤**會被測試抓到。`kernel` 有 Golden Case 與矩陣查表的驗收，
算錯就紅。

**回填不完整**不會。範本上某一格該填而空著，程式照樣跑完、驗證照樣通過、
測試照樣綠燈，因為沒有任何一個斷言在問「那一格是不是空的」。實際踩到三次：
表4 表頭六格整排沒填、表3 的建築密度值寫進標籤格把欄位名蓋掉、
土地利用現況整格沒動。三個都是交出去會被看出來的。

原因是範本沒有機器可讀的結構描述。哪一格是標籤、哪一格是值欄、哪些格
本案本來就該空白，全部只能靠實測與逐列判斷。

所以這裡提供兩層工具：

1. `audit_sheet()` 給人看。逐列印出範本內容與我們填的值，標記
   `[範本]` 沒動、`[填]` 從空白填入、`[改]` 覆寫了範本原有內容。
   改完格位對映或換範本之後跑一次，逐列掃過去。
2. `filled_coords()` 給測試看。回傳被動過的座標集合，讓測試釘住
   「應該填的格數與位置」，防止改動時安靜地少填一格。

## 三類標記怎麼讀

- `[填]` 是常態。
- `[改]` 要停下來想。合理的情況有兩種：把勾選符號 `□` 換成 `■`、
  把範本印的單位提示 `％` 換成帶百分比格式的數值。其他情況通常是
  把標籤蓋掉了，那就是 bug。
- `[範本]` 沒動過的格。要判斷是「本案沒有這一項」還是「漏填」。
  判斷依據是題目 PDF 的原始勘查表，不是猜的。
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterator, NamedTuple

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from .write import the_visible_sheet


class CellAudit(NamedTuple):
    """一格的盤點結果。"""

    coord: str
    row: int
    template: str
    """範本原本的內容，空字串代表空白"""
    output: str
    """產出檔案的內容"""

    @property
    def kind(self) -> str:
        if self.template == self.output:
            return "範本"
        if not self.template:
            return "填"
        return "改"

    def render(self, width: int = 24) -> str:
        if self.kind == "範本":
            return f"{self.coord}[範本]{self.template[:width]}"
        if self.kind == "填":
            return f"{self.coord}[填]{self.output[:width]}"
        return f"{self.coord}[改]{self.template[:14]}→{self.output[:width]}"


def _merge_anchors(ws: Worksheet) -> dict[tuple[int, int], str]:
    """(列, 欄) → 合併範圍的主格座標。沒合併的格不在裡面。"""
    out: dict[tuple[int, int], str] = {}
    for rng in ws.merged_cells.ranges:
        anchor = str(rng).split(":")[0]
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = anchor
    return out


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def audit_sheet(
    template: str | Path,
    output: str | Path,
    *,
    sheet_title: str | None = None,
    max_row: int | None = None,
) -> Iterator[CellAudit]:
    """逐格比對範本與產出。

    只回傳兩邊至少一邊有內容的格，而且跳過合併範圍的從格（那些格永遠是
    `None`，openpyxl 也不允許寫入）。

    `sheet_title` 是產出檔案裡的工作表名稱。表3 一個檔案有四張表
    （每個地價區段一張），所以要指定。範本那邊一律取唯一的可見工作表。
    """
    wbt = openpyxl.load_workbook(template)
    wst = the_visible_sheet(wbt)

    wbo = openpyxl.load_workbook(output)
    wso = wbo[sheet_title] if sheet_title else the_visible_sheet(wbo)

    anchors = _merge_anchors(wso)
    last = max_row or max(wst.max_row, wso.max_row)

    for r in range(1, last + 1):
        for c in range(1, max(wst.max_column, wso.max_column) + 1):
            coord = wso.cell(row=r, column=c).coordinate
            if anchors.get((r, c), coord) != coord:
                continue
            a = _text(wst.cell(row=r, column=c).value)
            b = _text(wso.cell(row=r, column=c).value)
            if not a and not b:
                continue
            yield CellAudit(coord=coord, row=r, template=a, output=b)


def filled_coords(
    template: str | Path,
    output: str | Path,
    *,
    sheet_title: str | None = None,
) -> set[str]:
    """被動過的座標集合（`[填]` 與 `[改]` 兩類）。

    測試拿這個釘住「填了哪些格」。少填一格集合就變小，測試會紅。
    """
    return {
        a.coord
        for a in audit_sheet(template, output, sheet_title=sheet_title)
        if a.kind != "範本"
    }


def format_report(
    template: str | Path,
    output: str | Path,
    *,
    sheet_title: str | None = None,
) -> str:
    """人看的逐列報告。"""
    by_row: dict[int, list[CellAudit]] = {}
    for a in audit_sheet(template, output, sheet_title=sheet_title):
        by_row.setdefault(a.row, []).append(a)

    lines = []
    for r in sorted(by_row):
        lines.append("R%-3d %s" % (r, "  ".join(x.render() for x in by_row[r])))

    counts = {"填": 0, "改": 0, "範本": 0}
    for row in by_row.values():
        for a in row:
            counts[a.kind] += 1
    lines.append("")
    lines.append(
        "小計：填入 %d 格、覆寫 %d 格、未動 %d 格"
        % (counts["填"], counts["改"], counts["範本"])
    )
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    """對一個產出目錄裡的三份書表逐列盤點。

        python -m xlsxform.audit --templates ../正式題目 --out <cli 的輸出目錄>

    先用 `python -m xlsxform.cli` 產出，再拿產出目錄餵給這支。
    """
    import argparse

    from . import layout
    from .pipeline import OUTPUT_STEM, find_template

    ap = argparse.ArgumentParser(description="回填完整性盤點：逐列比對範本與產出")
    ap.add_argument("--templates", type=Path, required=True, help="官方 xlsx 空白範本所在目錄")
    ap.add_argument("--out", type=Path, required=True, help="xlsxform.cli 的輸出目錄")
    ap.add_argument(
        "--segment",
        default=None,
        help="只看表3 的某個地價區段（例如 P001-00）。預設全部四張",
    )
    args = ap.parse_args(argv)

    def show(title: str, template: Path, output: Path, sheet: str | None = None) -> None:
        if not output.exists():
            print(f"（跳過 {title}：找不到 {output}）\n")
            return
        print("═" * 72)
        print(title)
        print("═" * 72)
        print(format_report(template, output, sheet_title=sheet))
        print()

    t3 = args.out / f"{OUTPUT_STEM['table3']}-filled.xlsx"
    if t3.exists():
        wb = openpyxl.load_workbook(t3)
        titles = [t for t in wb.sheetnames if t.startswith("表3 ")]
        if args.segment:
            titles = [t for t in titles if args.segment in t]
        for t in titles:
            show(f"表3 {t}", find_template(args.templates, "table3"), t3, sheet=t)

    show(
        "表5-1（定版）",
        find_template(args.templates, "table5"),
        args.out / f"{OUTPUT_STEM['table5']}-final.xlsx",
        sheet=layout.SHEET_TABLE5_1,
    )
    show(
        "表4（定版）",
        find_template(args.templates, "table4"),
        args.out / f"{OUTPUT_STEM['table4']}-final.xlsx",
        sheet=layout.SHEET_TABLE4,
    )

    print("讀法：[填] 是常態；[改] 要確認是勾選符號或單位提示，不是把標籤蓋掉；")
    print("      [範本] 要確認是「本案沒有這一項」而不是漏填，依據是題目原始勘查表。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
